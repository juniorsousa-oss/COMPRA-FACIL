"""Importação e exportação de listas de compras do Compra Fácil (.xlsx)."""

import io
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADERS = ("PRODUTO", "QUANTIDADE", "UNIDADE", "CATEGORIA", "ALTERNATIVA", "OBSERVAÇÃO")
ALIASES = {
    "name": ("produto", "nome", "nome do produto", "nome produto", "item", "descricao", "descrição"),
    "qty": ("quantidade", "qtd", "qtde", "qnt"),
    "unit": ("unidade", "und", "un"),
    "category": ("categoria", "grupo", "subgrupo"),
    "alt_name": ("alternativa", "produto alternativo", "opcao alternativa", "opção alternativa"),
    "observation": ("observacao", "observação", "obs"),
    "status": ("status", "situacao", "situação"),
}


def _text(value):
    return "" if value is None else str(value).strip()


def parse_excel(file_bytes, products, norm, find_product, num, parse_line):
    """Converte a planilha no mesmo formato de candidatos da importação por foto."""
    if not file_bytes or len(file_bytes) > 8 * 1024 * 1024:
        raise ValueError("Selecione um Excel .xlsx de até 8 MB.")
    aliases = {norm(alias): key for key, values in ALIASES.items() for alias in values}
    try:
        workbook = load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError("Não foi possível abrir o Excel. Envie um arquivo .xlsx válido.") from exc
    try:
        sheet = header = header_row = None
        for candidate_sheet in workbook.worksheets:
            for index, cells in enumerate(candidate_sheet.iter_rows(
                    min_row=1, max_row=min(candidate_sheet.max_row or 1, 12),
                    values_only=True), start=1):
                mapped = {}
                for position, value in enumerate(cells):
                    key = aliases.get(norm(value))
                    if key and key not in mapped:
                        mapped[key] = position
                if "name" in mapped:
                    sheet, header, header_row = candidate_sheet, mapped, index
                    break
            if sheet is not None:
                break
        if sheet is None:
            raise ValueError("Não encontrei a coluna PRODUTO. Utilize o modelo Excel disponível no aplicativo.")

        candidates, skipped, counted = [], [], 0
        for row_index, cells in enumerate(sheet.iter_rows(
                min_row=header_row + 1, values_only=True), start=header_row + 1):
            def cell(field):
                position = header.get(field)
                return cells[position] if position is not None and position < len(cells) else None

            name = _text(cell("name"))
            if not name:
                continue
            if norm(cell("status")) in {"ok", "confirmado", "comprado", "concluido", "feito"}:
                continue
            counted += 1
            if counted > 120:
                raise ValueError("Há mais de 120 produtos pendentes. Divida a planilha em arquivos menores.")
            raw_qty = cell("qty")
            qty = 1.0 if not _text(raw_qty) else num(raw_qty)
            if not (0 < qty < 100000):
                skipped.append(f"linha {row_index}: quantidade inválida")
                continue
            alternative = _text(cell("alt_name")) or None
            if not alternative and " ou " in name.lower():
                parsed = parse_line(name)
                if parsed and parsed.get("alt_name"):
                    name, alternative = parsed["name"], parsed["alt_name"]
            exact = find_product(products, name)
            alt_exact = find_product(products, alternative) if alternative else None
            category = (_text(cell("category"))
                        or (exact.get("categoria") if exact else "")
                        or "Mercearia")
            candidates.append({
                "raw": name, "name": name, "alt_name": alternative,
                "qty": float(qty), "unit": _text(cell("unit")) or "un.",
                "category": category, "observation": _text(cell("observation")),
                "confidence": None, "source": "Excel",
                # Nomes apenas parecidos serão revisados, sem troca automática.
                "suggested": exact.get("nome") if exact else None,
                "alt_suggested": alt_exact.get("nome") if alt_exact else None,
            })
        if not candidates:
            raise ValueError("Nenhum produto pendente com quantidade válida foi encontrado.")
        note = f"{len(candidates)} produto(s) encontrados na aba '{sheet.title}'."
        if skipped:
            note += f" {len(skipped)} linha(s) ignorada(s): " + "; ".join(skipped[:5])
        return {"candidates": candidates, "text": "", "ai_raw": "",
                "engine": "Excel", "ocr_confidence": None, "note": note}
    finally:
        workbook.close()


def build_excel(rows=(), template=False, num=float):
    """Gera o modelo preenchível ou exporta somente os itens pendentes."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Lista de compras"
    sheet.append(HEADERS)

    if template:
        info = workbook.create_sheet("Como preencher")
        for message in (
            "Preencha a aba Lista de compras, uma linha por produto.",
            "PRODUTO é obrigatório. QUANTIDADE é opcional e, quando vazia, equivale a 1.",
            "UNIDADE, CATEGORIA, ALTERNATIVA e OBSERVAÇÃO são opcionais.",
            "Exemplo: Arroz | 5 | kg | Mercearia.",
            "Você pode enviar uma foto ao ChatGPT e solicitar o preenchimento deste modelo.",
            "Os produtos só entram no Compra Fácil após a revisão e confirmação.",
        ):
            info.append([message])
        info.column_dimensions["A"].width = 96
    else:
        for item in rows:
            if item.get("confirmado"):
                continue

            def safe(value):
                content = _text(value)
                return "'" + content if content.startswith(("=", "+", "-", "@")) else content

            sheet.append([
                safe(item.get("nome_produto")),
                num(item.get("quantidade")) or 1,
                safe(item.get("unidade") or "un."),
                safe(item.get("categoria") or "Mercearia"),
                safe(item.get("produto_alternativo") or item.get("alternativa") or ""),
                "",
            ])

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = "A1:F" + str(max(sheet.max_row, 2))
    for index, width in enumerate((34, 16, 16, 26, 30, 40), start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="2563EB")
        cell.alignment = Alignment(horizontal="center")
    out = io.BytesIO()
    workbook.save(out)
    return out.getvalue()
