"""Testes de comparação financeira e leitura de planilha sem dados de produção."""
import io
import unittest

from openpyxl import Workbook
from receipt_audit import amount, compare, money, price_model, read_prices_excel


class ReceiptAuditTests(unittest.TestCase):
    def test_excel_preco_unitario_e_total(self):
        book = Workbook()
        sheet = book.active
        sheet.append(["PRODUTO", "QUANTIDADE", "PREÇO UNITÁRIO PAGO", "TOTAL PAGO"])
        sheet.append(["Achocolatado pronto", 5, 1.09, 5.45])
        stream = io.BytesIO()
        book.save(stream)
        parsed = read_prices_excel(stream.getvalue())
        self.assertEqual(len(parsed["lines"]), 1)
        self.assertEqual(parsed["lines"][0]["unit_price"], "1.09")
        self.assertEqual(parsed["lines"][0]["line_total"], "5.45")
        self.assertEqual(parsed["lines"][0]["qty"], "5")

    def test_modelo_preenchivel(self):
        model = price_model([{"nome_produto": "Arroz", "quantidade": 2}])
        self.assertTrue(model.startswith(b"PK"))
        book = __import__("openpyxl").load_workbook(io.BytesIO(model))
        self.assertEqual(book.active["A2"].value, "Arroz")
        self.assertIsNone(book.active["C2"].value)

    def test_total_correto(self):
        current = [{"id": 7, "nome_produto": "Achocolatado pronto",
                    "quantidade": 5, "preco_unitario": 1.09, "confirmado": True}]
        lines = [{"name": "Achocolatado pronto", "qty": "5",
                  "unit_price": "1.09", "line_total": "5.45"}]
        result = compare(lines, current, {0: "7"}, "5.45")
        self.assertEqual(result["results"][0]["status"], "Confere")
        self.assertEqual(result["results"][0]["total_app"], "R$ 5,45")
        self.assertEqual(result["results"][0]["total_doc"], "R$ 5,45")

    def test_quantidade_divergente(self):
        current = [{"id": 7, "nome_produto": "Café", "quantidade": 3,
                    "preco_unitario": 10, "confirmado": True}]
        lines = [{"name": "Café", "qty": "2", "unit_price": "10", "line_total": "20"}]
        report = compare(lines, current, {0: "7"})
        self.assertEqual(report["results"][0]["status"], "Quantidade diferente")

    def test_preco_divergente(self):
        current = [{"id": 7, "nome_produto": "Arroz", "quantidade": 2,
                    "preco_unitario": 12, "confirmado": True}]
        lines = [{"name": "Arroz", "qty": "2", "unit_price": "11", "line_total": "22"}]
        report = compare(lines, current, {0: "7"})
        self.assertEqual(report["results"][0]["status"], "Total divergente")
        self.assertEqual(report["results"][0]["diferenca"], "R$ -2,00")

    def test_quantidade_fracionaria(self):
        current = [{"id": 7, "nome_produto": "Carne", "quantidade": 0.345,
                    "preco_unitario": 20, "confirmado": True}]
        lines = [{"name": "Carne", "qty": "0.345",
                  "unit_price": "20", "line_total": "6.90"}]
        result = compare(lines, current, {0: "7"})
        self.assertEqual(result["results"][0]["status"], "Confere")
        self.assertEqual(amount(result["app_confirmed_total"]), amount("6.90"))

    def test_excel_sem_valores(self):
        book = Workbook()
        sheet = book.active
        sheet.append(["PRODUTO", "QUANTIDADE", "PREÇO UNITÁRIO PAGO"])
        sheet.append(["Café", 2, None])
        stream = io.BytesIO()
        book.save(stream)
        with self.assertRaisesRegex(ValueError, "Não foram encontrados"):
            read_prices_excel(stream.getvalue())

    def test_nenhum_vinculo_nao_confere(self):
        current = [{"id": 7, "nome_produto": "Leite", "quantidade": 1,
                    "preco_unitario": 5, "confirmado": True}]
        lines = [{"name": "Outra marca", "qty": "1", "unit_price": "5", "line_total": "5"}]
        result = compare(lines, current, {})
        self.assertEqual(result["results"][0]["status"], "Não encontrado no comprovante")
        self.assertEqual(result["extras"], ["Outra marca"])


if __name__ == "__main__":
    unittest.main()
